"""Reading small spreadsheets (class lists, rubrics) into rows of text.

Accepts ``.xlsx`` (first sheet) and ``.csv``. CSV files saved by Japanese
Excel are usually Shift-JIS, so UTF-8 is tried first and then cp932. The old
``.xls`` format cannot be read without extra libraries; the teacher is asked
to save as ``.xlsx`` instead.
"""

from __future__ import annotations

import csv
import io
import os
import unicodedata
from typing import List

from .errors import UserError

MAX_ROWS = 2000


def _cell_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return unicodedata.normalize("NFC", str(value)).strip()


def _trim(rows: List[List[str]]) -> List[List[str]]:
    rows = [r for r in rows if any(c for c in r)][:MAX_ROWS]
    width = max((len(r) for r in rows), default=0)
    # Drop columns that are empty in every row, then pad to a rectangle.
    keep = [i for i in range(width) if any(i < len(r) and r[i] for r in rows)]
    return [[r[i] if i < len(r) else "" for i in keep] for r in rows]


def decode_csv_bytes(raw: bytes) -> str:
    for encoding in ("utf-8-sig", "cp932"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise UserError("file_encoding_unknown")


def read_csv_bytes(raw: bytes) -> List[List[str]]:
    text = decode_csv_bytes(raw)
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    return _trim([[_cell_text(c) for c in row]
                  for row in csv.reader(io.StringIO(text), dialect)])


def read_xlsx_bytes(raw: bytes) -> List[List[str]]:
    from openpyxl import load_workbook
    try:
        wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    except Exception as exc:
        raise UserError("spreadsheet_unreadable") from exc
    try:
        ws = wb.worksheets[0]
        rows = []
        for row in ws.iter_rows(values_only=True):
            rows.append([_cell_text(c) for c in row])
            if len(rows) >= MAX_ROWS:
                break
    finally:
        wb.close()
    return _trim(rows)


def read_table(filename: str, raw: bytes) -> List[List[str]]:
    """Rows of cell text from an uploaded file's name and bytes."""
    ext = os.path.splitext(str(filename))[1].lower()
    if ext == ".xls":
        raise UserError("xls_not_supported")
    if ext in (".xlsx", ".xlsm"):
        return read_xlsx_bytes(raw)
    if ext in (".csv", ".txt", ".tsv"):
        return read_csv_bytes(raw)
    raise UserError("spreadsheet_type_unknown", file=os.path.basename(str(filename)))
