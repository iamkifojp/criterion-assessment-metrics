"""The master Excel gradebook: the teacher's permanent record of the exam.

Sheet 1, "Grades & Comments": one row per paper with every question's grade
and comment, the calculated result *and* the result on the report (with an
"adjusted" flag where the teacher overrode it), the overall comment, and
class average / highest / lowest rows at the bottom.

Sheet 2, "Questions & Rubrics": each question's boxes, marks or rubric and
choice group, then every rubric's levels and descriptors, so the file
explains itself years later.

All numbers are values worked out by grading_summary (no Excel formulas), so
the file always matches what CAM2 showed, including its rounding.
"""

from __future__ import annotations

import os

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from . import grading_summary, i18n, model, rubrics, slicer, storage
from .errors import UserError

HEADER_FILL = PatternFill("solid", fgColor="DCE6FA")
ADJUSTED_FILL = PatternFill("solid", fgColor="FFF2C2")
UNGRADED_FILL = PatternFill("solid", fgColor="FDE2E0")
SUMMARY_FILL = PatternFill("solid", fgColor="EEF1F6")
BOLD = Font(bold=True)
WRAP = Alignment(wrap_text=True, vertical="top")
TOP = Alignment(vertical="top")


def gradebook_path(session: storage.Session, config: dict) -> str:
    name = f"{slicer.safe_filename(config['exam_name'], 'Exam')}_Master_Grades.xlsx"
    return session.path(f"{storage.REPORTS_DIR}/{name}")


def build_workbook(config: dict, grades: dict, language: str) -> Workbook:
    from .workflow import display_name                     # avoid an import cycle
    t = lambda key, **kw: i18n.t(key, language, **kw)      # noqa: E731
    summary = grading_summary.class_summary(config, grades)
    questions = config.get("questions") or []
    rubric_list = config.get("rubrics") or []
    rubric_names = {r["id"]: r["name"] for r in rubric_list}
    mode = config.get("mode")
    na_text = t("grade.not_answered")
    # Per question: its grade columns (one per criterion it is assessed on), then its comment.
    q_rids = {q["id"]: ([None] if mode != "rubric"
                        else grading_summary.question_rubric_ids(config, q)) for q in questions}

    wb = Workbook()
    ws = wb.active
    ws.title = t("xlsx.sheet_grades")[:31]

    header = [t("xlsx.no"), t("xlsx.student"), t("xlsx.scan_file")]
    comment_cols = set()
    for q in questions:
        rids = q_rids[q["id"]]
        header += [q["label"] if len(rids) < 2 else t("xlsx.criterion_for", label=q["label"],
                                                        name=rubric_names[rid]) for rid in rids]
        header.append(t("xlsx.comment_for", label=q["label"]))
        comment_cols.add(len(header))
    if mode == "rubric":
        for r in rubric_list:
            header += [t("xlsx.average_for", name=r["name"]), t("xlsx.final_for", name=r["name"])]
    else:
        header += [t("xlsx.total"), t("xlsx.max"), t("xlsx.percent"), t("xlsx.report_total")]
    header += [t("xlsx.adjusted"), t("xlsx.overall_comment"), t("xlsx.complete")]
    ws.append(header)
    for cell in ws[1]:
        cell.font = BOLD
        cell.fill = HEADER_FILL
        cell.alignment = WRAP

    entries_all = grades.get("grades") or {}
    for i, (student, res) in enumerate(zip(config.get("students") or [], summary["students"]), 1):
        entries = entries_all.get(student["id"]) or {}
        row = [i, display_name(config, student), student.get("source") or student.get("file")]
        ungraded_cols = []
        for q in questions:
            entry = entries.get(q["id"]) or {}
            for rid in q_rids[q["id"]]:
                value = grading_summary.level_of(config, entry, q, rid)
                if value is None:
                    ungraded_cols.append(len(row) + 1)
                row.append(na_text if value == model.NOT_ANSWERED else value)
            row.append(entry.get("comment") or None)
        adjusted = False
        adjusted_cols = []
        if mode == "rubric":
            for r in rubric_list:
                rr = res["rubrics"][r["id"]]
                row += [rr["average"], rr["final"]]
                if rr["adjusted"]:
                    adjusted = True
                    adjusted_cols.append(len(row))
        else:
            m = res["marks"]
            row += [m["total"], m["max"], m["percent"], m["report_total"]]
            if m["adjusted"]:
                adjusted = True
                adjusted_cols.append(len(row))
        row += [t("xlsx.yes") if adjusted else None, res["comment"] or None,
                t("xlsx.yes") if res["complete"] else t("xlsx.no_answer")]
        ws.append(row)
        r_idx = ws.max_row
        for col in ungraded_cols:
            ws.cell(r_idx, col).fill = UNGRADED_FILL
        for col in adjusted_cols:
            ws.cell(r_idx, col).fill = ADJUSTED_FILL
        for cell in ws[r_idx]:
            cell.alignment = WRAP if isinstance(cell.value, str) and len(cell.value) > 20 else TOP

    # Class average / highest / lowest
    ws.append([])
    stats_rows = {key: [None, t(f"xlsx.{key}"), None] for key in ("average", "max_row", "min_row")}
    for q in questions:
        for rid in q_rids[q["id"]]:
            st = summary["questions"].get(grading_summary.slot_key(q["id"], rid))
            for key, field in (("average", "average"), ("max_row", "max"), ("min_row", "min")):
                stats_rows[key].append(st[field] if st else None)
        for key in stats_rows:
            stats_rows[key].append(None)
    if mode == "rubric":
        for r in rubric_list:
            st = summary["overall"].get(r["id"])
            for key, field in (("average", "average"), ("max_row", "max"), ("min_row", "min")):
                stats_rows[key] += [None, st[field] if st else None]
    else:
        st = summary["overall"].get("total")
        for key, field in (("average", "average"), ("max_row", "max"), ("min_row", "min")):
            stats_rows[key] += [st[field] if st else None, None, None, None]
    for key in ("average", "max_row", "min_row"):
        ws.append(stats_rows[key])
        for cell in ws[ws.max_row]:
            cell.fill = SUMMARY_FILL
            cell.font = BOLD

    widths = {1: 5, 2: 24, 3: 20}
    for col in range(1, ws.max_column + 1):
        title = str(ws.cell(1, col).value or "")
        is_comment = col in comment_cols
        ws.column_dimensions[get_column_letter(col)].width = widths.get(
            col, 36 if is_comment or col == ws.max_column - 1 else max(10, min(len(title) + 2, 22)))
    ws.freeze_panes = "C2"

    # Sheet 2: questions and rubrics --------------------------------------------------------
    ref = wb.create_sheet(t("xlsx.sheet_reference")[:31])
    ref.append([t("xlsx.question"), t("xlsx.boxes"),
                t("xlsx.max") if mode != "rubric" else t("xlsx.rubric"), t("xlsx.choice")])
    groups = {qid: g for g in config.get("choice_groups") or [] for qid in g["questions"]}
    labels = {q["id"]: q["label"] for q in questions}
    for q in questions:
        g = groups.get(q["id"])
        choice = (t("xlsx.choice_text", answer=g["answer"],
                    questions=", ".join(labels.get(x, x) for x in g["questions"])) if g else None)
        if mode == "rubric":
            third = ", ".join(rubric_names[rid] for rid in q_rids[q["id"]]) or None
        else:
            third = q.get("max")
        ref.append([q["label"], ", ".join(q.get("boxes") or []), third, choice])
    for cell in ref[1]:
        cell.font = BOLD
        cell.fill = HEADER_FILL
    if mode == "rubric" and rubric_list:
        ref.append([])
        ref.append([t("xlsx.rubric"), t("xlsx.levels"), t("xlsx.descriptor")])
        for cell in ref[ref.max_row]:
            cell.font = BOLD
            cell.fill = HEADER_FILL
        for r in rubric_list:
            for band in r["bands"]:
                ref.append([r["name"], band["levels"], band.get("descriptor") or None])
                ref.cell(ref.max_row, 3).alignment = WRAP
            ref.append([t("xlsx.rubric_max", name=r["name"], max=rubrics.rubric_max(r))])
    ref.column_dimensions["A"].width = 24
    ref.column_dimensions["B"].width = 34
    ref.column_dimensions["C"].width = 60
    ref.column_dimensions["D"].width = 34
    return wb


def export_gradebook(session: storage.Session, config: dict, grades: dict, language: str) -> str:
    path = gradebook_path(session, config)
    os.makedirs(os.path.dirname(path), exist_ok=True)
    wb = build_workbook(config, grades, language)
    tmp = path + ".tmp"
    try:
        wb.save(tmp)
        os.replace(tmp, path)
    except PermissionError as exc:      # the gradebook is open in Excel
        raise UserError("report_file_open", file=os.path.basename(path)) from exc
    finally:
        if os.path.exists(tmp):
            os.remove(tmp)
    return path
