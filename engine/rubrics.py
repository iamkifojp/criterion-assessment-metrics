"""Teacher-made rubrics.

CAM2 ships no rubric text. A rubric is a name plus a list of bands; each band
covers one level ("3") or a range of levels ("5–6") and has a descriptor::

    {"id": "r1", "name": "Criterion A",
     "bands": [{"levels": "0", "descriptor": "..."},
               {"levels": "1–2", "descriptor": "..."}]}

New rubrics start with the familiar MYP rows 0, 1–2, 3–4, 5–6, 7–8, but the
teacher may change, add or remove rows freely. The highest level in any band
becomes the rubric's maximum. When grading, every level inside a band gets its
own button, so a "5–6" band offers 5 and 6.
"""

from __future__ import annotations

import io
import re
import unicodedata
from typing import List, Optional, Tuple

from .errors import UserError

DEFAULT_BANDS = ("0", "1–2", "3–4", "5–6", "7–8")
MAX_LEVEL = 100

# NFKC turns full-width digits and "－"/"～" into ASCII; "〜" and "ー" stay.
_LEVELS_RE = re.compile(r"^\s*(\d+)\s*(?:[-–—~〜ー]\s*(\d+))?\s*$")


def parse_levels(text) -> Tuple[int, int]:
    """'3' -> (3, 3); '5-6', '5–6', '5〜6', '５－６' -> (5, 6)."""
    m = _LEVELS_RE.match(unicodedata.normalize("NFKC", str(text or "")))
    if not m:
        raise UserError("bad_levels", levels=str(text or ""))
    lo = int(m.group(1))
    hi = int(m.group(2)) if m.group(2) else lo
    if lo > hi or hi > MAX_LEVEL:
        raise UserError("bad_levels", levels=str(text))
    return lo, hi


def levels_label(lo: int, hi: int) -> str:
    return str(lo) if lo == hi else f"{lo}–{hi}"


def new_rubric(rubric_id: str, name: str = "") -> dict:
    return {"id": rubric_id, "name": name,
            "bands": [{"levels": b, "descriptor": ""} for b in DEFAULT_BANDS]}


def normalize_rubric(raw: dict) -> dict:
    """Validate a rubric from the UI or a file and tidy its level labels.

    Bands keep the teacher's order. Raises UserError on bad or overlapping
    levels, or a rubric with no bands.
    """
    name = str(raw.get("name") or "").strip()
    bands = []
    for band in raw.get("bands") or []:
        levels = str(band.get("levels") or "").strip()
        descriptor = str(band.get("descriptor") or "").strip()
        if not levels and not descriptor:
            continue                        # an empty row the teacher left
        lo, hi = parse_levels(levels)
        bands.append({"levels": levels_label(lo, hi), "min": lo, "max": hi,
                      "descriptor": descriptor})
    if not bands:
        raise UserError("rubric_no_bands", rubric=name)
    spans = sorted((b["min"], b["max"], b["levels"]) for b in bands)
    for (lo1, hi1, a), (lo2, _hi2, b) in zip(spans, spans[1:]):
        if lo2 <= hi1:
            raise UserError("rubric_levels_overlap", rubric=name, a=a, b=b)
    return {"id": str(raw.get("id") or ""), "name": name, "bands": bands}


def rubric_max(rubric: dict) -> int:
    return max(int(b["max"]) for b in rubric["bands"])


def band_for_level(rubric: dict, level: int) -> Optional[dict]:
    for band in rubric["bands"]:
        if int(band["min"]) <= level <= int(band["max"]):
            return band
    return None


def check_level(rubric: dict, level) -> int:
    if isinstance(level, bool) or not isinstance(level, int):
        raise UserError("level_not_in_rubric", level=level, rubric=rubric.get("name", ""))
    if band_for_level(rubric, level) is None:
        raise UserError("level_not_in_rubric", level=level, rubric=rubric.get("name", ""))
    return level


# --- Excel / CSV import and template ---------------------------------------------------

_HEADER_KEYWORDS = (
    # Checked in this order, so "Level descriptor" counts as a descriptor column.
    ("descriptor", ("descriptor", "description", "記述", "説明", "内容")),
    ("criterion", ("criterion", "criteria", "rubric", "基準", "観点", "ルーブリック")),
    ("levels", ("level", "band", "grade", "score", "mark", "レベル", "バンド", "段階", "点")),
)

TEMPLATE_HEADERS = {
    "en": ("Criterion", "Levels", "Descriptor"),
    "ja": ("基準", "レベル", "記述"),
}


def _header_columns(row: List[str]) -> Optional[dict]:
    found = {}
    for i, cell in enumerate(row):
        text = cell.casefold()
        for kind, words in _HEADER_KEYWORDS:
            if kind not in found and any(w in text for w in words):
                found[kind] = i
                break
    return found if {"levels", "descriptor"} <= set(found) else None


def import_rubrics(rows: List[List[str]]) -> List[dict]:
    """Rubrics from spreadsheet rows (columns Criterion | Levels | Descriptor).

    A header row is recognised in English or Japanese; without one, three
    columns are read as criterion/levels/descriptor and two as
    levels/descriptor. A blank criterion cell continues the one above
    (as merged cells read).
    """
    if not rows:
        raise UserError("rubric_file_empty")
    cols = _header_columns(rows[0])
    body = rows[1:] if cols else rows
    if not cols:
        width = max(len(r) for r in rows)
        if width >= 3:
            cols = {"criterion": 0, "levels": 1, "descriptor": 2}
        elif width == 2:
            cols = {"levels": 0, "descriptor": 1}
        else:
            raise UserError("rubric_columns_unknown")

    def cell(row, kind):
        i = cols.get(kind)
        return row[i] if i is not None and i < len(row) else ""

    rubrics: List[dict] = []
    current = None
    for row in body:
        name = cell(row, "criterion")
        if current is None or (name and name != current["name"]):
            current = {"id": f"r{len(rubrics) + 1}", "name": name or "Rubric 1", "bands": []}
            rubrics.append(current)
        if cell(row, "levels") or cell(row, "descriptor"):
            current["bands"].append({"levels": cell(row, "levels"),
                                     "descriptor": cell(row, "descriptor")})
    return [normalize_rubric(r) for r in rubrics if r["bands"]]


def template_xlsx(language: str = "en") -> bytes:
    """A blank rubric spreadsheet the teacher can fill in."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font

    headers = TEMPLATE_HEADERS.get(language, TEMPLATE_HEADERS["en"])
    example = "基準A" if language == "ja" else "Criterion A"
    wb = Workbook()
    ws = wb.active
    ws.title = "Rubric"
    ws.append(list(headers))
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for levels in DEFAULT_BANDS:
        ws.append([example, levels, ""])
    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 10
    ws.column_dimensions["C"].width = 80
    for row in ws.iter_rows(min_row=2):
        row[2].alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
