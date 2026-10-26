"""Tests for the Excel master's Classroom Entry tab layout (2026-10-27).

The tab lists a **Grades** block (one column per folder assignment, or per
criterion when an assignment assessed several) followed by a **Comments** block
(one column per assignment). Draft assignments carry feedback but no grade, so
their grade column reads ``NA``. Pure stdlib ``unittest`` with stubbed app
state; no Streamlit run, no database. Run:

    python -m unittest tests.test_classroom_entry_layout
"""

import io
import os
import sys
import unittest
from types import SimpleNamespace

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import app  # noqa: E402  (path shim must precede the import)


def _score(assignment, crit, value, comment="", valid=True):
    return SimpleNamespace(assignment=assignment,
                           criterion=SimpleNamespace(value=crit),
                           value=value, comment=comment, is_valid=valid)


class ClassroomEntryLayoutTests(unittest.TestCase):
    def setUp(self):
        self._orig = {n: getattr(app, n) for n in
                      ("st", "gb", "_classroom_folder_assignment_names",
                       "all_scores")}
        roster = [
            {"key": "1", "name": "Aoki Daichi", "first": "Daichi"},
            {"key": "2", "name": "Baba Chiaki", "first": "Chiaki"},
        ]
        app.st = SimpleNamespace(session_state={"roster": roster,
                                                "active_class": "8.1"})
        self.students = [SimpleNamespace(student_id=e["key"], name=e["name"])
                         for e in roster]
        assignments = [
            SimpleNamespace(name="Unit 1", class_name="8.1", criteria=["A"],
                            is_draft=False, draft_feedback={}),
            SimpleNamespace(name="Draft", class_name="8.1", criteria=["A"],
                            is_draft=True,
                            draft_feedback={"1": {"comment": "Good start."}}),
            SimpleNamespace(name="Unit 2", class_name="8.1", criteria=["B", "A"],
                            is_draft=False, draft_feedback={}),
            SimpleNamespace(name="Other class", class_name="8.2", criteria=["A"],
                            is_draft=False, draft_feedback={}),
        ]
        app.gb = lambda: SimpleNamespace(assignments=assignments)
        app._classroom_folder_assignment_names = lambda: ["Unit 1", "Draft", "Unit 2"]
        by_id = {s.student_id: s for s in self.students}
        scores = [
            (by_id["1"], _score("Unit 1", "A", 5, "Nice shading.")),
            (by_id["2"], _score("Unit 1", "A", 0, "wrong file", valid=False)),
            (by_id["2"], _score("Unit 1", "A", 6, "Strong work.")),
            (by_id["1"], _score("Unit 2", "A", 4, "Same text.")),
            (by_id["1"], _score("Unit 2", "B", 7, "Same text.")),
            (by_id["2"], _score("Unit 2", "B", 3, "Only B.")),
            (by_id["1"], _score("Other class", "A", 8, "ignored")),
        ]
        app.all_scores = lambda: scores

    def tearDown(self):
        for n, v in self._orig.items():
            setattr(app, n, v)

    def _build(self):
        from openpyxl import Workbook, load_workbook
        wb = Workbook()
        app._append_classroom_entry_sheet(wb, list(self.students))
        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        return load_workbook(buf)["Classroom Entry"]

    def _row(self, ws, r):
        return [ws.cell(row=r, column=c).value for c in range(1, ws.max_column + 1)]

    def test_headers_grades_block_then_comments_block(self):
        ws = self._build()
        self.assertEqual(ws.cell(row=1, column=3).value, "Grades")
        self.assertEqual(ws.cell(row=1, column=7).value, "Comments")
        self.assertEqual(self._row(ws, 2)[2:], [
            "Unit 1", "Draft", "Unit 2 – Crit A", "Unit 2 – Crit B",
            "Unit 1", "Draft", "Unit 2"])
        merged = {str(r) for r in ws.merged_cells.ranges}
        self.assertIn("C1:F1", merged)
        self.assertIn("G1:I1", merged)

    def test_rows_hold_grades_then_comments(self):
        ws = self._build()
        # Latin first-name order: Chiaki (id 2) before Daichi (id 1).
        self.assertEqual(self._row(ws, 3), [
            "Baba Chiaki", "2", 6, "NA", None, 3,
            "Strong work.", None, "Only B."])
        self.assertEqual(self._row(ws, 4), [
            "Aoki Daichi", "1", 5, "NA", 4, 7,
            "Nice shading.", "Good start.", "Same text."])

    def test_no_assignments_leaves_fixed_columns_only(self):
        app._classroom_folder_assignment_names = lambda: []
        ws = self._build()
        self.assertEqual(ws.max_column, 2)
        self.assertEqual(ws.freeze_panes, "C3")


class FolderAssignmentFilterTests(unittest.TestCase):
    """Drafts now qualify; other 0-criterion events and exams still don't."""

    def setUp(self):
        self._orig = {n: getattr(app, n) for n in ("st", "gb", "assignment_table")}
        app.st = SimpleNamespace(session_state={"active_class": "8.1"})
        mk = lambda name, src="x.csv": SimpleNamespace(
            name=name, class_name="8.1", source_file=src)
        app.gb = lambda: SimpleNamespace(assignments=[
            mk("Graded"), mk("Draft"), mk("Formative"), mk("Exam"),
            mk("Local", src=""), mk("Crit D (Reflection)")])
        row = lambda name, formative=False, draft=False, exam=False: {
            "name": name, "is_formative": formative, "is_draft": draft,
            "is_exam": exam, "folder_ref": ""}
        app.assignment_table = lambda: [
            row("Graded"), row("Draft", formative=True, draft=True),
            row("Formative", formative=True), row("Exam", exam=True),
            row("Local"), row("Crit D (Reflection)")]

    def tearDown(self):
        for n, v in self._orig.items():
            setattr(app, n, v)

    def test_drafts_included_others_filtered(self):
        self.assertEqual(app._classroom_folder_assignment_names(),
                         ["Graded", "Draft"])


if __name__ == "__main__":
    unittest.main()
