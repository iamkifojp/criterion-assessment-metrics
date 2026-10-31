"""Student PDF reports and the Excel gradebook.

Run:  python -m unittest tests.test_reports
"""

import os
import tempfile
import unittest
from datetime import date
from unittest import mock

import pymupdf as fitz
from openpyxl import load_workbook

from engine import gradebook_xlsx, grading_summary as gs, model, reports, storage, workflow
from engine.errors import UserError
from engine.slicer import Layout
from tests.helpers import make_pdf

RED = (255, 0, 0)


class ReportTestCase(unittest.TestCase):
    mode = "marks"

    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.root = self.tmp.name
        env = mock.patch.dict(os.environ, {"CAM2_APP_DATA": os.path.join(self.root, "appdata")})
        env.start()
        self.addCleanup(env.stop)
        scans = os.path.join(self.root, "scans")
        os.makedirs(scans)
        for name in ("Tanaka Hana.pdf", "Sato Ken.pdf", "scan_003.pdf", "scan_004.pdf"):
            make_pdf(os.path.join(scans, name), pages=2, fills={1: [(Layout("A4"), 2, 3, RED)]})
        self.session, self.config, self.grades = storage.create_session(
            os.path.join(self.root, "session"), "Y8: Science / Test")
        workflow.import_scans(self.session, self.config, {"type": "folder", "path": scans})
        workflow.set_mode(self.config, self.mode)
        if self.mode == "rubric":
            workflow.set_rubrics(self.config, [{"name": "Criterion A", "bands": [
                {"levels": "0", "descriptor": "Nothing yet."},
                {"levels": "1-2", "descriptor": "States facts."},
                {"levels": "3-4", "descriptor": "Explains with examples."}]}])
            questions = [{"id": "q1", "label": "Q1", "boxes": ["page1!C4"]},
                         {"id": "q2", "label": "Q2", "boxes": ["page2!A1:B2"]}]
        else:
            questions = [{"id": "q1", "label": "Q1", "boxes": ["page1!C4"], "max": 5},
                         {"id": "q2", "label": "Q2", "boxes": ["page2!A1:B2"], "max": 3},
                         {"id": "q3", "label": "Q3", "boxes": ["page2!C3"], "max": 3}]
        workflow.set_questions(self.config, {"name_box": ["page1!A1:B2"], "questions": questions,
                                             "choice_groups": [{"questions": ["q2", "q3"], "answer": 1}]
                                             if self.mode == "marks" else []})
        workflow.slice_session(self.session, self.config)
        # Papers sort as: Sato Ken, scan_003, scan_004, Tanaka Hana
        workflow.set_roster(self.config, [{"id": "n1", "name": "Hana Tanaka"},
                                          {"id": "n2", "name": "Ken Sato"}])

    def tearDown(self):
        self.tmp.cleanup()

    def text_of(self, path):
        with fitz.open(path) as doc:
            return "\n".join(page.get_text() for page in doc)


class MarksReportTests(ReportTestCase):
    def setUp(self):
        super().setUp()
        g = lambda sid, qid, v, c=None: gs.set_grade(self.config, self.grades, sid, qid, v, c)  # noqa: E731
        g("s001", "q1", 4, "Good use of evidence.")
        g("s001", "q2", 2)
        g("s001", "q3", 3)                         # choice: q3 counts, q2 doesn't
        g("s002", "q1", model.NOT_ANSWERED)
        gs.set_report_total(self.grades, "s001", 8)
        gs.set_overall_comment(self.grades, "s001", "A strong paper. 次も頑張りましょう。")

    def test_report_content(self):
        path = os.path.join(self.root, "r.pdf")
        reports.build_student_pdf(self.session, self.config, self.grades, self.config["students"][0],
                                  "en", path, day=date(2026, 10, 30))
        text = self.text_of(path)
        for expected in ("Y8: Science / Test", "Ken Sato", "30 October 2026", "8 / 8",      # reported total
                         "Good use of evidence.", "A strong paper.", "次も頑張りましょう", "not counted"):
            self.assertIn(expected, text)
        self.assertLess(os.path.getsize(path), 400_000)
        with fitz.open(path) as doc:          # one image per answer; no handwritten name crop
            self.assertEqual(sum(len(page.get_images()) for page in doc), 3)

    def test_teacher_picked_choice(self):
        self.config["choice_groups"][0]["pick"] = "teacher"
        path = os.path.join(self.root, "pick.pdf")

        def text():
            reports.build_student_pdf(self.session, self.config, self.grades, self.config["students"][0],
                                      "en", path, day=date(2026, 10, 30))
            return " ".join(self.text_of(path).split())

        self.assertNotIn("not counted", text())             # both answered: waiting for a pick
        gs.set_choice(self.config, self.grades, "s001", self.config["choice_groups"][0]["id"], ["q2"])
        self.assertIn("your teacher chose which answers count", text())

    def test_japanese_report(self):
        path = os.path.join(self.root, "ja.pdf")
        reports.build_student_pdf(self.session, self.config, self.grades, self.config["students"][1],
                                  "ja", path, day=date(2026, 10, 30))
        text = self.text_of(path)
        for expected in ("2026年10月30日", "合計", "無回答", "未採点"):
            self.assertIn(expected, text)

    def test_export_names_and_errors(self):
        self.config["students"][2]["name"] = "Ken Sato"          # a second student with the same name
        real = reports.build_student_pdf

        def flaky(session, config, grades, student, language, out_path, day=None):
            if student["id"] == "s004":
                raise PermissionError("open in a PDF viewer")
            return real(session, config, grades, student, language, out_path, day)

        with mock.patch.object(reports, "build_student_pdf", side_effect=flaky):
            result = reports.export_student_pdfs(self.session, self.config, self.grades, "en")
        # s001 Sato Ken (class list), s002 scan_003 (no name), s003 typed "Ken Sato", s004 fails
        self.assertEqual(result["files"], ["Ken Sato_Y8_ Science _ Test.pdf",
                                           "UNMATCHED_scan_003.pdf",
                                           "Ken Sato_Y8_ Science _ Test (2).pdf"])
        self.assertEqual([e["student"] for e in result["errors"]], ["s004"])
        self.assertTrue(all(os.path.isfile(os.path.join(result["folder"], f)) for f in result["files"]))
        # Making them again replaces the old files rather than adding to them
        again = reports.export_student_pdfs(self.session, self.config, self.grades, "en")
        self.assertEqual(sorted(os.listdir(again["folder"])), sorted(again["files"]))
        self.assertIn("Hana Tanaka_Y8_ Science _ Test.pdf", again["files"])

    def test_gradebook(self):
        path = gradebook_xlsx.export_gradebook(self.session, self.config, self.grades, "en")
        self.assertTrue(path.endswith("Y8_ Science _ Test_Master_Grades.xlsx"))
        wb = load_workbook(path)
        ws = wb["Grades & Comments"]
        header = [c.value for c in ws[1]]
        self.assertEqual(header[:5], ["No.", "Student", "Scan file", "Q1", "Q1 comment"])
        row = {h: c.value for h, c in zip(header, ws[2])}
        self.assertEqual((row["Student"], row["Q1"], row["Q1 comment"]), ("Ken Sato", 4, "Good use of evidence."))
        self.assertEqual((row["Calculated total"], row["Out of"], row["Total on report"]), (7, 8, 8))
        self.assertEqual(row["Adjusted by teacher"], "Yes")
        self.assertEqual(ws.cell(2, header.index("Total on report") + 1).fill.fgColor.rgb[-6:], "FFF2C2")
        row2 = {h: c.value for h, c in zip(header, ws[3])}
        self.assertEqual(row2["Q1"], "Not answered")
        labels = [ws.cell(r, 2).value for r in range(1, ws.max_row + 1)]
        self.assertIn("Class average", labels)
        ref = wb["Questions & Rubrics"]
        self.assertEqual(ref.cell(3, 4).value, "Answer 1 of Q2, Q3")

    def test_gradebook_open_in_excel(self):
        with mock.patch("openpyxl.Workbook.save", side_effect=PermissionError("locked")):
            with self.assertRaises(UserError) as ctx:
                gradebook_xlsx.export_gradebook(self.session, self.config, self.grades, "en")
        self.assertEqual(ctx.exception.code, "report_file_open")


class RubricReportTests(ReportTestCase):
    mode = "rubric"

    def test_descriptor_final_level_and_japanese_gradebook(self):
        gs.set_grade(self.config, self.grades, "s001", "q1", 3, "Nice example.")
        gs.set_grade(self.config, self.grades, "s001", "q2", 4)
        gs.set_final_level(self.config, self.grades, "s001", "r1", 3)
        path = os.path.join(self.root, "r.pdf")
        reports.build_student_pdf(self.session, self.config, self.grades, self.config["students"][0], "en", path)
        text = self.text_of(path)
        for expected in ("Criterion A", "3 / 4", "Explains with examples.", "Nice example.",
                         "States facts.", "Nothing yet."):           # the whole rubric, once
            self.assertIn(expected, text)
        self.assertEqual(text.count("Explains with examples."), 1)
        self.assertNotIn("· 3–4", text)
        wb = load_workbook(gradebook_xlsx.export_gradebook(self.session, self.config, self.grades, "ja"))
        self.assertEqual(wb.sheetnames, ["成績・コメント一覧", "設問配点・ルーブリック基準"])
        ws = wb["成績・コメント一覧"]
        header = [c.value for c in ws[1]]
        row = {h: c.value for h, c in zip(header, ws[2])}
        self.assertEqual((row["Criterion A 平均"], row["Criterion A 最終レベル"], row["先生が調整"]),
                         (3.5, 3, "はい"))


class MultiCriterionReportTests(ReportTestCase):
    mode = "rubric"

    def setUp(self):
        super().setUp()
        workflow.set_rubrics(self.config, self.config["rubrics"] + [{"name": "Criterion B", "bands": [
            {"levels": "0"}, {"levels": "1-8", "descriptor": "Designs a method."}]}])
        for q in self.config["questions"]:
            q["rubric_ids"] = ["r1", "r2"] if q["id"] == "q2" else ["r1"]
        g = lambda qid, v, rid=None: gs.set_grade(self.config, self.grades, "s001", qid, v, rubric_id=rid)  # noqa: E731
        g("q1", 2)
        g("q2", 4, "r1")
        g("q2", 7, "r2")
        gs.set_grade(self.config, self.grades, "s002", "q2", model.NOT_ANSWERED)

    def test_report_shows_each_criterion(self):
        path = os.path.join(self.root, "r.pdf")
        reports.build_student_pdf(self.session, self.config, self.grades, self.config["students"][0], "en", path)
        text = self.text_of(path)
        self.assertIn("Criterion A 4 / 4 · Criterion B 7 / 8", " ".join(text.split()))
        self.assertIn("Designs a method.", text)
        reports.build_student_pdf(self.session, self.config, self.grades, self.config["students"][1], "en", path)
        text = self.text_of(path)
        self.assertNotIn("Criterion A Not answered", " ".join(text.split()))

    def test_gradebook_has_a_column_per_criterion(self):
        wb = load_workbook(gradebook_xlsx.export_gradebook(self.session, self.config, self.grades, "en"))
        ws = wb["Grades & Comments"]
        header = [c.value for c in ws[1]]
        self.assertEqual(header[3:8], ["Q1", "Q1 comment", "Q2 · Criterion A", "Q2 · Criterion B", "Q2 comment"])
        row = {h: c.value for h, c in zip(header, ws[2])}
        self.assertEqual((row["Q2 · Criterion A"], row["Q2 · Criterion B"]), (4, 7))
        self.assertEqual((row["Criterion A average"], row["Criterion B average"]), (3.0, 7.0))
        row2 = {h: c.value for h, c in zip(header, ws[3])}
        self.assertEqual(row2["Q2 · Criterion B"], "Not answered")
        averages = next(r for r in ws.iter_rows(values_only=True) if r[1] == "Class average")
        self.assertEqual(averages[header.index("Q2 · Criterion B")], 3.5)
        ref = wb["Questions & Rubrics"]
        self.assertEqual(ref.cell(3, 3).value, "Criterion A, Criterion B")


if __name__ == "__main__":
    unittest.main()
