"""Fictional cross-group reporting, detached exports and durable references."""
from copy import deepcopy
from datetime import datetime
import io
import os
from pathlib import Path
import tempfile
from types import SimpleNamespace
import unittest
from unittest import mock
from uuid import uuid4
import zipfile

import app
from engine.models import Assignment, CriterionScore, Criterion, Student, UnitPlan
from engine.persistence import (save_database, load_database, validate_database_payload,
                                persistent_content_fingerprint, serialize_gradebook)
from engine.reporting import (ReportScope, ReportMember, resolve_members,
                              definition_issues, content_fingerprint)


class ReportingTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.env = mock.patch.dict(os.environ, {'CAM_DB_PATH': self.tmp.name})
        self.env.start()
        self.addCleanup(self.env.stop)
        self.original_st = app.st
        app.st = SimpleNamespace(session_state={'db_loaded': True})
        self.addCleanup(setattr, app, 'st', self.original_st)
        with mock.patch.object(app, 'load_prefs', return_value=dict(app.DEFAULT_PREFS)):
            app.init_state()
        self.ss = app.st.session_state
        self.term = app.TERMS[0]
        self.ss.update(active_class='G3', active_term=self.term,
            classes=[dict(name=g, grade='Year 8', myp_year='3', subject='Arts')
                     for g in ('G1', 'G2', 'G3')],
            rosters={g: [dict(key=sid, name=name, first=first, email=f'{sid}@example.test')]
                     for g, sid, name, first in [('G1', 's1', 'Baba Emi', 'Emi'),
                        ('G2', 's2', 'Aoki Ken', 'Ken'), ('G3', 's3', 'Chiba Akira', 'Akira')]},
            unit_plans={g: UnitPlan(unit_title=f'Unit {g}') for g in ('G1', 'G2', 'G3')})
        self.ss['rosters']['G1'].append(dict(key='empty', name='Iida Bob', first='Bob', email='empty@example.test'))
        self.ss['rosters']['G1'].append(dict(key='draft', name='Sato Alice', first='Alice', email='draft@example.test'))
        for group, rows in self.ss['rosters'].items():
            for row in rows:
                app.gb().students[row['key']] = Student(row['key'], row['name'])
        for i, group in enumerate(('G1', 'G2', 'G3'), 1):
            app.gb().assignments.append(Assignment('Shared title', ['A'], class_name=group,
                term=self.term, ingested_at=datetime(2026, 4, i)))
            app.gb().students[f's{i}'].add_score(CriterionScore(Criterion.A, i+4,
                datetime(2026, 4, i), assignment='Shared title', include_in_report=False))
        app.gb().assignments.extend([
            Assignment('Extra G2', ['B'], class_name='G2', term=self.term, ingested_at=datetime(2026, 5, 1)),
            Assignment('Awaiting', ['C'], class_name='G1', term=self.term, folder_ref='fictional', ingested_at=datetime(2026, 5, 2)),
            Assignment('Draft', ['D'], class_name='G1', term=self.term, is_draft=True,
                       draft_feedback={'draft': {'comment': 'Qualitative only'}}, ingested_at=datetime(2026, 5, 3)),
            Assignment('Old task', ['D'], class_name='G1', term=app.TERMS[1], ingested_at=datetime(2026, 7, 1))])
        app.gb().students['s1'].add_score(CriterionScore(Criterion.D, 8, datetime(2026, 7, 1), assignment='Old task'))
        self.ss['comments_by_term'] = {self.term: {'s1': '  Saved exactly.\nSecond line.  ', 's2': 'Other saved comment'}}
        self.ss['teacher_remarks']['s1'] = ' Teacher remark '
        self.ss['final_override']['s2'] = {'A': 8}
        self.ss['excused_flags']['s2||Extra G2'] = True
        app.ensure_class_context()
        app.ensure_term_context()
        self.definition = dict(id=str(uuid4()), name='Rocky', sort_mode='manual', members=[
            dict(student_id='s2', source_class='G2'), dict(student_id='s1', source_class='G1')])
        self.ss['reporting_classes'] = [self.definition]

    def scope(self, definition=None):
        definition = definition or self.definition
        members, errors = resolve_members(definition, self.ss, app.gb(), app.sort_roster)
        self.assertEqual(errors, [])
        return ReportScope(definition['id'], definition['name'], self.term, members, True)

    def export(self, scope=None):
        state, book = app.reporting_snapshot()
        return app.ReportingExport(scope or self.scope(), state, book)

    def test_roundtrip_validation_and_legacy_fingerprint(self):
        payload = app.build_session_payload()
        path = os.path.join(self.tmp.name, 'test.json')
        save_database(path, app.gb(), payload)
        loaded = load_database(path)
        self.assertEqual(loaded['session']['reporting_classes'], [self.definition])
        old = {k: v for k, v in payload.items() if k != 'reporting_classes'}
        self.assertEqual(persistent_content_fingerprint(app.gb(), old),
                         persistent_content_fingerprint(app.gb(), dict(old, reporting_classes=[])))
        self.ss['reporting_classes'] = [dict(self.definition, name='Changed')]
        app.restore_session(loaded['session'])
        self.assertEqual(self.ss['reporting_classes'], [self.definition])
        app.restore_session(old)
        self.assertEqual(self.ss['reporting_classes'], [])
        for mutation in (dict(self.definition, sort_mode='invalid'),
                         dict(self.definition, members=self.definition['members'] * 2),
                         dict(self.definition, name=' '), dict(self.definition, id='not-a-uuid')):
            self.assertTrue(definition_issues([mutation]))
            issues = validate_database_payload(dict(version=1, gradebook=serialize_gradebook(app.gb()),
                                                       session={'reporting_classes': [mutation]}))
            self.assertTrue(issues)
            self.assertNotIn('s1', str(issues))
        self.assertTrue(definition_issues([self.definition, dict(self.definition, id=str(uuid4()), name='rocky')]))

    def test_own_group_grades_missing_and_context_equivalence(self):
        report = self.export()
        for ctx, student in report.members:
            source = ctx.st.session_state['active_class']
            default = self.export(ReportScope('teaching', source, self.term,
                tuple(ReportMember(r['key'], source) for r in self.ss['rosters'][source])))
            default_ctx = default.contexts[source]
            self.assertEqual(ctx.student_term_grades(student), default_ctx.student_term_grades(default_ctx.gb().students[student.student_id]))
        first = report.contexts['G1']
        self.assertEqual(first.aggregate_with_policy(first.gb().students['s1'], 'A').rounded_band, 5)
        self.assertIsNone(first.aggregate_with_policy(first.gb().students['s1'], 'B'))
        self.assertIsNone(first.aggregate_with_policy(first.gb().students['s1'], 'C'))
        self.assertIsNone(first.aggregate_with_policy(first.gb().students['s1'], 'D'))
        self.assertEqual(first.aggregate_with_policy(first.gb().students['empty'], 'A').rounded_band, 0)
        self.assertIsNone(first.aggregate_with_policy(first.gb().students['draft'], 'D'))
        self.assertIsNone(report.contexts['G2'].aggregate_with_policy(report.contexts['G2'].gb().students['s2'], 'B'))
        self.ss['excused_flags'].clear()
        missing = self.export().contexts['G2']
        self.assertEqual(missing.aggregate_with_policy(missing.gb().students['s2'], 'B').rounded_band, 0)

    def test_all_exports_order_source_provenance_comments_and_no_mutation(self):
        from openpyxl import load_workbook
        from docx import Document
        before = deepcopy((app.build_session_payload(), app.gb(), self.ss['active_class'], self.ss['roster'], self.ss['unit_plan']))
        report = self.export()
        wb = load_workbook(io.BytesIO(report.excel()))
        self.assertEqual(wb.sheetnames, ['Final Suggestions', 'Raw Scores', 'Assignments'])
        self.assertEqual([r[0] for r in list(wb['Final Suggestions'].values)[1:]], ['s2', 's1'])
        self.assertEqual(wb['Final Suggestions']['C2'].value, 8)
        self.assertEqual(wb['Final Suggestions']['J1'].fill.fgColor.rgb[-6:], 'B3554D')
        self.assertEqual(wb['Final Suggestions']['J2'].value, 'G2')
        shared = [r for r in list(wb['Assignments'].values)[6:] if r[0] == 'Shared title']
        self.assertEqual({r[-1] for r in shared}, {'G1', 'G2'})
        self.assertEqual({r[4] for r in shared}, {5, 6})
        raw = list(wb['Raw Scores'].values)[1:]
        self.assertIn('Old task', [r[1] for r in raw])
        self.assertNotIn('s3', [r[0] for r in raw])
        text = '\n'.join(p.text for p in Document(io.BytesIO(report.docx())).paragraphs)
        self.assertLess(text.index('Aoki Ken'), text.index('Baba Emi'))
        self.assertIn('Reporting class: Rocky', text)
        self.assertIn('Unit G1', text)
        self.assertIn('Unit G2', text)
        self.assertNotIn('Unit G3', text)
        comments = [p.text for p in Document(io.BytesIO(report.comments())).paragraphs]
        self.assertIn('  Saved exactly.\nSecond line.  ', comments)
        self.assertIn(' Teacher remark ', comments)
        packed, skipped = report.zip()
        self.assertEqual(skipped, [])
        self.assertEqual(zipfile.ZipFile(io.BytesIO(packed)).namelist(), ['s2@example.test.docx', 's1@example.test.docx'])
        single_doc = Document(io.BytesIO(report.docx('s1')))
        marks = [cell.text for row in single_doc.tables[0].rows for cell in row.cells]
        self.assertIn('Apr 01', marks)
        self.assertNotIn('Extra G2', marks)
        self.assertNotIn('Old task', marks)
        single = '\n'.join(p.text for p in single_doc.paragraphs)
        self.assertIn('Baba Emi', single)
        self.assertNotIn('Aoki Ken', single)
        self.assertEqual(before, (app.build_session_payload(), app.gb(), self.ss['active_class'], self.ss['roster'], self.ss['unit_plan']))
        # Fail midway after a source helper has accessed its context.
        ctx = report.contexts['G1']
        ctx._student_docx.__globals__['_trend_png'] = mock.Mock(side_effect=RuntimeError('injected'))
        with self.assertRaises(RuntimeError):
            report.docx()
        self.assertEqual(before, (app.build_session_payload(), app.gb(), self.ss['active_class'], self.ss['roster'], self.ss['unit_plan']))

    def test_orders_unresolved_ambiguity_and_metadata(self):
        original = deepcopy(self.ss['rosters'])
        for mode in ('manual', 'last_first', 'first_last', 'gojuon', 'email'):
            definition = dict(self.definition, sort_mode=mode)
            scope = self.scope(definition)
            expected = ['s2', 's1'] if mode == 'manual' else [r['key'] for r in app.sort_roster([
                self.ss['rosters']['G2'][0], self.ss['rosters']['G1'][0]], mode)]
            self.assertEqual([m.student_id for m in scope.members], expected)
        self.assertEqual(self.ss['rosters'], original)
        self.ss['archived_students']['G1'] = [self.ss['rosters']['G1'][0]]
        _, errors = resolve_members(self.definition, self.ss, app.gb(), app.sort_roster)
        self.assertTrue(errors)
        self.ss['archived_students'].clear()
        self.ss['classes'][1]['subject'] = 'Music'
        _, errors = resolve_members(self.definition, self.ss, app.gb(), app.sort_roster)
        self.assertTrue(any('conflicting' in e for e in errors))
        self.ss['classes'][1]['subject'] = 'Arts'
        self.ss['rosters']['G2'].append(self.ss['rosters']['G1'][0])
        _, errors = resolve_members(self.definition, self.ss, app.gb(), app.sort_roster)
        self.assertTrue(any('ambiguous' in e for e in errors))

    def test_fingerprint_detects_same_count_swaps_and_content_edits(self):
        baseline = self.export().fingerprint
        self.definition['members'].reverse()
        self.assertNotEqual(baseline, self.export().fingerprint)
        self.definition['members'].reverse()
        member = self.definition['members'][1]
        member['student_id'] = 'empty'
        self.assertNotEqual(baseline, self.export().fingerprint)
        member['student_id'] = 's1'
        changes = [lambda: self.definition.update(name='Andes'),
                   lambda: self.ss['comments_by_term'][self.term].update(s1='edited'),
                   lambda: self.ss['report_cfg'].update(show_effort=True),
                   lambda: setattr(app.gb().assignments[-1], 'is_draft', True),
                   lambda: self.ss['rosters']['G1'][0].update(email='changed@example.test')]
        for change in changes:
            before = self.export().fingerprint
            change()
            self.assertNotEqual(before, self.export().fingerprint)

    def test_default_excel_keeps_classroom_entry(self):
        from openpyxl import load_workbook
        scope = ReportScope('teaching', 'G1', self.term,
                            tuple(ReportMember(r['key'], 'G1') for r in self.ss['rosters']['G1']))
        wb = load_workbook(io.BytesIO(self.export(scope).excel()))
        self.assertIn('Classroom Entry', wb.sheetnames)
        self.assertEqual(wb['Classroom Entry']['A3'].value, 'Sato Alice')

    def test_sort_order_applies_to_each_output_without_rewriting_manual_order(self):
        from openpyxl import load_workbook
        from docx import Document
        manual = deepcopy(self.definition['members'])
        for mode in ('manual', 'last_first', 'first_last', 'gojuon', 'email'):
            self.definition['sort_mode'] = mode
            report = self.export()
            expected = [m.student_id for m in report.scope.members]
            wb = load_workbook(io.BytesIO(report.excel()))
            self.assertEqual([r[0] for r in list(wb['Final Suggestions'].values)[1:]], expected)
            packed, _ = report.zip()
            self.assertEqual(zipfile.ZipFile(io.BytesIO(packed)).namelist(),
                             [sid+'@example.test.docx' for sid in expected])
            for data in (report.docx(), report.comments()):
                text = '\n'.join(p.text for p in Document(io.BytesIO(data)).paragraphs)
                labels = [app.gb().students[sid].name for sid in expected]
                self.assertLess(text.index(labels[0]), text.index(labels[1]))
            self.assertEqual(self.definition['members'], manual)

    def test_exam_missing_policy_uses_source_group_and_draft_only_members_resolve(self):
        from engine.models import ExamResult
        from openpyxl import load_workbook
        app.gb().assignments.append(Assignment('Exam G1', ['B'], class_name='G1', term=self.term,
            is_exam=True, max_total=10, ingested_at=datetime(2026, 5, 4)))
        student = app.gb().students['empty']
        student.add_score(CriterionScore(Criterion.B, 6, datetime(2026, 5, 4), assignment='Exam G1'))
        student.exam_results['Exam G1'] = ExamResult('Exam G1', total=8, max_total=10)
        report = self.export()
        ctx = report.contexts['G1']
        # The selected subset has no submitted exam. Banding in its source
        # group still establishes this as genuine missing work, with zero.
        self.assertEqual(ctx.aggregate_with_policy(ctx.gb().students['s1'], 'B').rounded_band, 0)
        wb = load_workbook(io.BytesIO(report.excel()))
        self.assertEqual(wb['Final Suggestions']['D3'].value, '0*')
        exam_rows = [r for r in list(wb['Assignments'].values)[6:] if r[0] == 'Exam G1']
        self.assertEqual(exam_rows[0][3], 0)  # subset analytics remains a subset
        self.ss['rosters']['G1'] = [r for r in self.ss['rosters']['G1'] if r['key'] != 'draft']
        self.definition['members'].append(dict(student_id='draft', source_class='G1'))
        self.assertIn('draft', [s.student_id for _, s in self.export().members])

    def test_mailmerge_checks_visibility_and_draft_marks(self):
        from docx import Document
        self.ss['report_cfg'].update(show_effort=True, show_myp_grade=True, show_school_grade=True)
        report = self.export()
        packed, _ = report.zip()
        archive = zipfile.ZipFile(io.BytesIO(packed))
        doc = Document(io.BytesIO(archive.read('s1@example.test.docx')))
        cells = [cell.text for table in doc.tables for row in table.rows for cell in row.cells]
        self.assertNotIn('Draft', cells)
        self.assertNotIn('Effort / English Use', cells)
        self.assertNotIn('School Grade', cells)
        self.assertIn('MYP Grade', cells)
        self.ss['rosters']['G1'][0]['email'] = 's2@example.test'
        self.assertEqual(self.export().zip()[1], [('Baba Emi', 'duplicate email')])
        self.ss['rosters']['G1'][0]['email'] = 'bad/address@example.test'
        self.assertEqual(self.export().zip()[1], [('Baba Emi', 'email unusable as filename')])
        self.ss['rosters']['G1'][0]['email'] = ''
        self.assertEqual(self.export().zip()[1], [('Baba Emi', 'no email on roster')])

    def test_term_restore_recovery_replacement_and_source_delete(self):
        from engine.persistence import (capture_database_snapshot, database_write_token,
            replace_database_checked, write_conflict_recovery)
        saved = deepcopy(self.ss['reporting_classes'])
        backup = app.build_term_backup(self.term)
        with mock.patch.object(app, 'persist'), mock.patch.object(app, '_pre_restore_backup', return_value=''):
            app.restore_term_backup(backup)
        self.assertEqual(self.ss['reporting_classes'], saved)
        path = os.path.join(self.tmp.name, 'lifecycle.json')
        save_database(path, app.gb(), app.build_session_payload())
        observed = capture_database_snapshot(path)
        recovered = write_conflict_recovery(path, app.gb(), app.build_session_payload(),
                            database_write_token(observed), observed)
        self.assertEqual(load_database(recovered.path)['session']['reporting_classes'], saved)
        replace_database_checked(path, app.gb(), app.build_session_payload())
        self.assertEqual(load_database(path)['session']['reporting_classes'], saved)
        with mock.patch.object(app, 'persist'), mock.patch.object(app, '_mark_teacher_input_deleted'):
            app.delete_class('G1')
        self.assertEqual(self.ss['reporting_classes'], saved)
        _, errors = resolve_members(self.definition, self.ss, app.gb(), app.sort_roster)
        self.assertTrue(any('unresolved' in error for error in errors))

    def test_isolated_ui_selection_manager_save_cancel_and_search(self):
        # Run only the tray/dialog in Streamlit's in-process UI harness. All
        # persistence, prefs, mirrors, watch roots and workspace paths are temp.
        from streamlit.testing.v1 import AppTest
        from engine.persistence import capture_database_snapshot, database_write_token
        path = os.path.join(self.tmp.name, 'acm_database.json')
        save_database(path, app.gb(), app.build_session_payload())
        self.ss['db_write_token'] = database_write_token(capture_database_snapshot(path))
        self.ss['prefs'].update(setup_done=True, db_custom_path=self.tmp.name)
        initial = deepcopy(self.ss)
        script = """
import streamlit as st
import app
app.st = st
assert app.db_folder() == SANDBOX
if st.session_state.get('test_manager'):
    app.reporting_classes_dialog()
else:
    app.render_tray()
""".replace('SANDBOX', repr(self.tmp.name))
        with mock.patch.object(app, 'PREFS_PATH', os.path.join(self.tmp.name, 'prefs.json')), \
             mock.patch.object(app, 'GRADING_WORKSPACE_DIR', os.path.join(self.tmp.name, 'workspace')), \
             mock.patch.object(app, '_mirror_classes_to_cloud'), \
             mock.patch.object(app, 'sync_from_cloud', side_effect=AssertionError('Cloud access forbidden')):
            at = AppTest.from_string(script)
            for key, value in initial.items():
                at.session_state[key] = value
            at.run(timeout=20)
            self.assertFalse(at.exception)
            before_selection = Path(path).read_bytes()
            at.selectbox(key='report_scope_id').set_value(self.definition['id']).run()
            self.assertFalse(at.exception)
            self.assertEqual(at.session_state['active_class'], 'G3')
            next(b for b in at.button if b.label.startswith('Build Excel')).click().run()
            self.assertFalse(at.exception)
            self.assertIn('xlsx', at.session_state['export_ready'])
            for label in ('Build report-card pack', 'Build mail-merge pack', 'Build class comments', 'Build report —'):
                next(b for b in at.button if b.label.startswith(label)).click().run()
                self.assertFalse(at.exception)
            self.assertEqual(set(at.session_state['export_ready']), {'xlsx', 'pack', 'mailmerge', 'comments', 'single'})
            at.session_state['final_override']['s2']['A'] = 7
            at.run()
            self.assertFalse(at.session_state['export_ready'])
            self.assertEqual(Path(path).read_bytes(), before_selection)
            at.session_state['test_manager'] = True
            at.run()
            at.text_input(key='report_edit_new_name').set_value('New fictional list').run()
            at.multiselect(key='report_edit_new_sources').select('G1').run()
            next(b for b in at.button if b.label == 'Select all — G1').click().run()
            at.text_input(key='report_edit_new_query').set_value('Baba').run()
            at.text_input(key='report_edit_new_query').set_value('Iida').run()
            next(b for b in at.button if b.label == 'Save reporting class').click().run()
            self.assertFalse(at.exception)
            self.assertFalse(at.error, [e.value for e in at.error])
            self.assertTrue(any(d['name'] == 'New fictional list' for d in at.session_state['reporting_classes']), str(at.session_state.filtered_state))
            definition = next(d for d in at.session_state['reporting_classes'] if d['name'] == 'New fictional list')
            self.assertEqual({m['student_id'] for m in definition['members']}, {'s1', 'empty', 'draft'})
            at.selectbox(key='report_manager_id').set_value(definition['id']).run()
            at.text_input(key=f"report_edit_{definition['id']}_name").set_value('Unsaved rename').run()
            before = deepcopy(at.session_state['reporting_classes'])
            next(b for b in at.button if b.label == 'Cancel').click().run()
            self.assertFalse(at.exception)
            self.assertEqual(at.session_state['reporting_classes'], before)
            self.assertEqual(load_database(path)['session']['reporting_classes'], before)

    def test_rename_and_wipe_lifecycle(self):
        with mock.patch.object(app, 'persist'), mock.patch.object(app, 'class_data_dir', side_effect=lambda cls: os.path.join(self.tmp.name, cls)):
            self.assertTrue(app.rename_class('G1', 'Renamed'))
            self.assertEqual(self.definition['members'][1]['source_class'], 'Renamed')
            app.wipe_database_full()
            self.assertEqual(self.ss['reporting_classes'], [])


if __name__ == '__main__':
    unittest.main()
